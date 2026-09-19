from pathlib import Path

from openpyxl import Workbook
import pytest

ROOT = Path(__file__).resolve().parents[4]
WORKBOOK = ROOT / 'docs/Global_Success_Khung_Nghe_Noi_3_Level_v3.xlsx'


def test_imported_unit_matches_committed_vocabulary_and_patterns():
    from luna_tutor.curriculum.excel_importer import import_workbook
    from luna_tutor.curriculum.loader import load_unit
    imported = import_workbook(WORKBOOK, grade=5, unit=1)
    committed = load_unit(ROOT / 'curriculum/grade-05/unit-01')
    assert imported.vocabulary_ids() == committed.vocabulary_ids()
    assert {p.id for p in imported.patterns} == {p.id for p in committed.patterns}
    assert len(imported.vocabulary) == 19
    assert imported.lessons[2].objective_ids == imported.lessons[0].objective_ids + imported.lessons[1].objective_ids
    assert not any('.lesson03.' in o.id for o in imported.objectives)
    assert all('!D' in v.source.section or '!F' in v.source.section or '!H' in v.source.section for v in imported.vocabulary)


@pytest.fixture
def bounded_workbook(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Lớp 5 (3 Level)'
    sheet.append(['Chủ đề (Unit)', 'Lesson', 'Ngữ âm', 'LEVEL 1 – Từ vựng', 'LEVEL 1 – Mẫu câu', 'LEVEL 2 – Từ vựng', 'LEVEL 2 – Mẫu câu', 'LEVEL 3 – Từ vựng', 'LEVEL 3 – Mẫu câu'])
    sheet.append(['Unit 1: First', 'Lesson 1', None, 'city', "I live in the ___.", 'hobby', "What's your hobby? – My hobby is ___.", 'calm'])
    sheet.append([None, 'Lesson 2', None, 'pink', "What's your favourite ___? – It's ___."])
    sheet.append([None, 'Lesson 3', None, 'Ôn tập từ vựng Lesson 1 + 2 + làm Project', 'Ôn tập 2 mẫu câu của Unit'])
    sheet.append(['Unit 2: Second', 'Lesson 1', None, 'house'])
    sheet.append([None, 'Lesson 2', None, 'flat'])
    sheet.merge_cells('A2:A4')
    sheet.merge_cells('H2:H4')
    path = tmp_path / 'curriculum.xlsx'
    workbook.save(path)
    return path


def test_merged_and_blank_unit_cells_do_not_duplicate_content(bounded_workbook):
    from luna_tutor.curriculum.excel_importer import import_workbook
    imported = import_workbook(bounded_workbook, 5, 1)
    assert imported.vocabulary_ids() == {'city', 'pink', 'hobby', 'calm'}
    assert len(imported.vocabulary) == 4
    assert imported.lessons[2].objective_ids == [
        'unit01.lesson01.vocabulary.city', 'unit01.lesson01.pattern.live_in',
        'unit01.lesson02.vocabulary.pink', 'unit01.lesson02.pattern.favourite',
    ]


def test_forward_fill_stops_at_next_unit(bounded_workbook):
    from luna_tutor.curriculum.excel_importer import import_workbook
    imported = import_workbook(bounded_workbook, 5, 2)
    assert imported.vocabulary_ids() == {'house', 'flat'}
    assert imported.patterns == []
    assert not any('.level02.' in o.id or '.level03.' in o.id for o in imported.objectives)


def test_unknown_unit_and_grade_fail_explicitly(bounded_workbook):
    from luna_tutor.curriculum.excel_importer import import_workbook
    with pytest.raises(ValueError, match='Unit 99'):
        import_workbook(bounded_workbook, 5, 99)
    with pytest.raises(ValueError, match='Grade 4'):
        import_workbook(bounded_workbook, 4, 1)
